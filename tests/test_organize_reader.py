# tests/test_organize_reader.py
"""Tests for fda.organize.reader."""

from __future__ import annotations

import os
import time
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest


def _summarizer_response(type_label: str = "doc", summary: str = "x"):
    """Return a `complete()` response payload that looks like our skill output.

    Uses json.dumps so the result is valid JSON (repr() would emit single
    quotes which json.loads rejects).
    """
    import json
    return json.dumps({"type_label": type_label, "summary": summary})


@pytest.fixture
def fake_backend():
    """Backend whose .complete returns a JSON summary payload."""
    backend = MagicMock()
    backend.complete.return_value = _summarizer_response()
    return backend


@pytest.fixture
def workspace(tmp_path):
    root = tmp_path / "ws"
    root.mkdir()
    return root


@pytest.fixture
def logger(tmp_path):
    from fda.organize._logger import OrganizeLogger
    return OrganizeLogger(log_path=tmp_path / "r.log", target_basename="ws")


class TestPathIdAssignment:
    def test_ids_assigned_in_path_sorted_order(self, workspace, fake_backend, logger):
        from fda.organize import reader

        (workspace / "b.txt").write_text("b")
        (workspace / "a.txt").write_text("a")
        (workspace / "c.txt").write_text("c")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        ids_paths = [(e.path_id, Path(e.path).name) for e in catalog.entries]
        assert ids_paths == [("f000", "a.txt"), ("f001", "b.txt"), ("f002", "c.txt")]
        # Happy-path assertion: at least one entry has summary_failed=False
        # so we know the JSON parsing path actually succeeded.
        assert any(not e.summary_failed for e in catalog.entries), (
            "expected at least one successful summary on happy path"
        )
        assert all(e.summary == "x" for e in catalog.entries if not e.summary_failed)

    def test_two_runs_same_tree_produce_same_ids(self, workspace, fake_backend, logger):
        from fda.organize import reader

        (workspace / "a.txt").write_text("a")
        (workspace / "b.txt").write_text("b")
        c1 = reader.read(workspace, backend=fake_backend, logger=logger)
        c2 = reader.read(workspace, backend=fake_backend, logger=logger)
        assert [e.path_id for e in c1.entries] == [e.path_id for e in c2.entries]
        assert [e.path for e in c1.entries] == [e.path for e in c2.entries]


class TestJunkFlagged:
    def test_ds_store_marked_is_junk_no_llm_call(self, workspace, fake_backend, logger):
        from fda.organize import reader

        (workspace / ".DS_Store").write_bytes(b"\x00")
        (workspace / "a.txt").write_text("a")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        junk = [e for e in catalog.entries if e.is_junk]
        assert [Path(e.path).name for e in junk] == [".DS_Store"]
        # Backend was called only for the non-junk file
        called_paths = [
            call.kwargs.get("messages", [{}])[0].get("content", "")
            for call in fake_backend.complete.call_args_list
        ]
        assert all(".DS_Store" not in c for c in called_paths)
        assert fake_backend.complete.call_count == 1


class TestGitWorktreeSkipped:
    def test_repo_root_and_internals_excluded(self, workspace, fake_backend, logger):
        from fda.organize import reader

        repo = workspace / "repo"
        repo.mkdir()
        (repo / ".git").mkdir()
        (repo / "main.py").write_text("# git-tracked")
        (workspace / "a.txt").write_text("a")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        names = [Path(e.path).name for e in catalog.entries]
        assert "main.py" not in names
        assert "a.txt" in names
        assert str(repo) in catalog.git_repos_skipped


class TestSymlinkSkipped:
    def test_file_and_dir_symlinks_excluded(self, workspace, fake_backend, logger):
        from fda.organize import reader

        outside = workspace.parent / "outside"
        outside.mkdir()
        (outside / "leak.txt").write_text("leak")
        (workspace / "a.txt").write_text("a")
        link_file = workspace / "link.txt"
        link_dir = workspace / "linkdir"
        try:
            os.symlink(outside / "leak.txt", link_file)
            os.symlink(outside, link_dir)
        except OSError:
            pytest.skip("symlinks unsupported on this platform")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        names = [Path(e.path).name for e in catalog.entries]
        assert "leak.txt" not in names
        assert "link.txt" not in names
        assert "a.txt" in names


class TestTextCap:
    def test_64k_truncation_marker_in_prompt(self, workspace, fake_backend, logger):
        from fda.organize import reader

        # 2 MB text file
        big = "z" * (2 * 1024 * 1024)
        (workspace / "big.txt").write_text(big)
        reader.read(workspace, backend=fake_backend, logger=logger)
        # The first .complete call's user message should contain
        # the [TRUNCATED at 64KB] marker.
        first = fake_backend.complete.call_args_list[0]
        user = first.kwargs.get("messages")[0]["content"]
        assert "[TRUNCATED at 64KB]" in user
        # The text content sent must be <= 64KB + small overhead
        assert len(user) < 80 * 1024


class TestPerFileFailure:
    def test_backend_error_records_failed_summary(self, workspace, logger):
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.side_effect = RuntimeError("boom")
        (workspace / "a.txt").write_text("a")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.summary == ""

    def test_unparseable_response_records_failed(self, workspace, logger):
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.return_value = "not-json"
        (workspace / "a.txt").write_text("a")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True


class TestPerFileTimeout:
    def test_backend_timeout_marks_failed(self, workspace, logger):
        """Per-file timeout enforced at the backend HTTP layer."""
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.side_effect = TimeoutError("backend HTTP timeout")
        (workspace / "a.txt").write_text("a")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True


class TestDeadline:
    def test_total_cap_marks_unstarted_failed(self, workspace, logger, monkeypatch):
        from fda.organize import reader

        # 12 files. With a deadline of 0.2s, the backend mock honors `timeout`
        # so calls dispatched late get a 0-ish timeout and raise immediately.
        for i in range(12):
            (workspace / f"f{i}.txt").write_text("x")
        monkeypatch.setattr(reader, "READER_TOTAL_TIMEOUT_SECONDS", 0.2)
        monkeypatch.setattr(reader, "READER_WORKER_COUNT", 2)

        backend = MagicMock()

        def slow(*, timeout=None, **kwargs):
            # Simulate a 0.5s HTTP call honoring the request timeout.
            if timeout is None or timeout >= 0.5:
                time.sleep(0.5)
                return _summarizer_response()
            time.sleep(max(0.0, timeout))
            raise TimeoutError("HTTP request deadline exceeded")

        backend.complete.side_effect = slow
        catalog = reader.read(workspace, backend=backend, logger=logger)
        failed = [e for e in catalog.entries if e.summary_failed]
        assert len(failed) >= 1, "expected at least one entry to fail by deadline"


class TestExtractorRegistryUsed:
    def test_register_runtime_extension(self, workspace, fake_backend, logger):
        from fda.organize import reader, _extractors
        from fda.organize.models import ExtractionResult

        (workspace / "a.foo").write_bytes(b"binary")

        def fake(_path):
            return ExtractionResult(text="hello-foo", status="ok")

        with patch.dict(_extractors.EXTRACTORS, {".foo": fake}, clear=False):
            reader.read(workspace, backend=fake_backend, logger=logger)
        first = fake_backend.complete.call_args_list[0]
        assert "hello-foo" in first.kwargs.get("messages")[0]["content"]


class TestNoDenylistShortCircuit:
    def test_register_for_exe_works(self, workspace, fake_backend, logger):
        """A format previously in _BINARY_EXTS is no longer denylisted."""
        from fda.organize import reader, _extractors
        from fda.organize.models import ExtractionResult

        (workspace / "a.exe").write_bytes(b"\x00\x01")

        def fake(_path):
            return ExtractionResult(text="exe-text", status="ok")

        with patch.dict(_extractors.EXTRACTORS, {".exe": fake}, clear=False):
            reader.read(workspace, backend=fake_backend, logger=logger)
        first = fake_backend.complete.call_args_list[0]
        assert "exe-text" in first.kwargs.get("messages")[0]["content"]


class TestNoExtractorStub:
    def test_unknown_extension_sends_stub(self, workspace, fake_backend, logger):
        from fda.organize import reader

        (workspace / "a.zzz").write_bytes(b"\x00")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.extract_status == "no_extractor"
        # Backend was called and given a stub describing the file.
        first = fake_backend.complete.call_args_list[0]
        msg = first.kwargs.get("messages")[0]["content"]
        assert "no_extractor" in msg or "binary" in msg or "stub" in msg


class TestSkillLoaderUnit:
    def test_load_skill_parses_frontmatter(self, tmp_path):
        from fda.organize import _skills

        skill_dir = tmp_path / "demo"
        skill_dir.mkdir()
        (skill_dir / "SKILL.md").write_text(
            "---\n"
            "name: demo\n"
            "description: a demo\n"
            "model: claude-haiku-4-5-20251001\n"
            "---\n"
            "Body here.\n"
        )
        cfg = _skills.load_skill(skill_dir)
        assert cfg.name == "demo"
        assert cfg.model == "claude-haiku-4-5-20251001"
        assert "Body here." in cfg.body

    def test_missing_model_raises(self, tmp_path):
        from fda.organize import _skills

        skill_dir = tmp_path / "demo"
        skill_dir.mkdir()
        (skill_dir / "SKILL.md").write_text("---\nname: demo\n---\nbody\n")
        with pytest.raises(ValueError):
            _skills.load_skill(skill_dir)


# ---------------------------------------------------------------------------
# F1 — TOCTOU guard: worker exception yields failed entry, not abort
# ---------------------------------------------------------------------------


class TestWorkerExceptionDoesNotAbortRun:
    def test_summarize_one_raise_yields_failed_entry(
        self, workspace, fake_backend, logger, monkeypatch
    ):
        """If _summarize_one raises (e.g., file vanished mid-walk), the run
        completes with a synthesized failed entry — it doesn't abort."""
        from fda.organize import reader

        (workspace / "a.txt").write_text("a")
        (workspace / "b.txt").write_text("b")

        original = reader._summarize_one
        call_count = {"n": 0}

        def flaky(path, **kwargs):
            call_count["n"] += 1
            if call_count["n"] == 1:
                raise OSError("file vanished mid-walk")
            return original(path, **kwargs)

        monkeypatch.setattr(reader, "_summarize_one", flaky)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        # 2 entries returned, one of them is failed
        assert len(catalog.entries) == 2
        failed = [e for e in catalog.entries if e.summary_failed]
        assert len(failed) == 1


# ---------------------------------------------------------------------------
# F2 — elapsed_ms reflects worker wall time, not as_completed overhead
# ---------------------------------------------------------------------------


class TestElapsedMsCovered:
    def test_elapsed_ms_is_nonzero_for_slow_call(self, workspace, tmp_path):
        """elapsed_ms should reflect worker wall time, not as_completed overhead."""
        import re as _re
        from fda.organize._logger import OrganizeLogger
        from fda.organize import reader

        log_path = tmp_path / "r.log"
        log = OrganizeLogger(log_path=log_path, target_basename="ws")

        (workspace / "a.txt").write_text("a")
        backend = MagicMock()

        def slow(*args, **kwargs):
            time.sleep(0.05)
            return _summarizer_response()

        backend.complete.side_effect = slow
        reader.read(workspace, backend=backend, logger=log)
        log.close()
        lines = log_path.read_text().splitlines()
        done_lines = [l for l in lines if "READER_FILE_DONE" in l]
        assert done_lines, "expected a READER_FILE_DONE event"
        match = _re.search(r"elapsed_ms=(\d+)", done_lines[0])
        assert match, "elapsed_ms field missing"
        assert int(match.group(1)) >= 50, f"elapsed_ms looks wrong: {done_lines[0]}"


# ---------------------------------------------------------------------------
# F3 — APITimeoutError → built-in TimeoutError rewrap
# ---------------------------------------------------------------------------


class TestApiBackendTimeoutRewrap:
    def test_apitimeouterror_rewrapped_as_builtin_timeouterror(self, monkeypatch):
        """The API backend translates anthropic.APITimeoutError to built-in
        TimeoutError so Reader's TimeoutError catch can classify it."""
        try:
            from anthropic import APITimeoutError
        except ImportError:
            pytest.skip("anthropic SDK not installed")

        from fda.claude_backend import AnthropicAPIBackend

        backend = AnthropicAPIBackend.__new__(AnthropicAPIBackend)

        # Build a stub client whose with_options(...) returns a stub whose
        # messages.create raises APITimeoutError.
        class _Stub:
            class messages:
                @staticmethod
                def create(**kwargs):
                    raise APITimeoutError(request=None)

        class _Client:
            def with_options(self, **kw):
                return _Stub()

            class messages:
                @staticmethod
                def create(**kwargs):
                    raise APITimeoutError(request=None)

        backend._client = _Client()
        with pytest.raises(TimeoutError):
            backend.complete(
                system="x",
                messages=[{"role": "user", "content": "hi"}],
                timeout=1.0,
            )


# ---------------------------------------------------------------------------
# F5 — deadline-skipped entries report extract_status="failed", not "ok"
# ---------------------------------------------------------------------------


class TestDeadlineExtractStatus:
    def test_deadline_skipped_entries_have_failed_extract_status(
        self, workspace, logger, monkeypatch
    ):
        """When a file is skipped because the deadline expired before the
        worker picked it up, the synthesized fail entry must report
        extract_status='failed', not 'ok'."""
        from fda.organize import reader

        for i in range(8):
            (workspace / f"f{i}.txt").write_text("x")

        # Force every worker pickup to find remaining <= 0.
        monkeypatch.setattr(reader, "READER_TOTAL_TIMEOUT_SECONDS", -1.0)

        backend = MagicMock()
        backend.complete.return_value = _summarizer_response()
        catalog = reader.read(workspace, backend=backend, logger=logger)
        # Every entry should be a deadline failure with extract_status="failed"
        for e in catalog.entries:
            assert e.summary_failed
            assert e.extract_status == "failed", (
                f"expected extract_status='failed' for deadline-skipped {e.path}, "
                f"got {e.extract_status!r}"
            )
        # Backend was not called for any file (all skipped pre-pickup).
        assert backend.complete.call_count == 0


# ---------------------------------------------------------------------------
# F6 — detail field present in failure log
# ---------------------------------------------------------------------------


class TestFailureDetailInLog:
    def test_unparseable_failure_logs_detail(self, workspace, tmp_path):
        """Unparseable JSON failures must surface a 'detail' field in the
        READER_FILE_FAIL log line so post-mortem can distinguish parser
        errors from backend errors."""
        from fda.organize._logger import OrganizeLogger
        from fda.organize import reader

        log_path = tmp_path / "r.log"
        log = OrganizeLogger(log_path=log_path, target_basename="ws")

        (workspace / "a.txt").write_text("a")
        backend = MagicMock()
        backend.complete.return_value = "not-json"
        reader.read(workspace, backend=backend, logger=log)
        log.close()
        lines = log_path.read_text().splitlines()
        fail_lines = [l for l in lines if "READER_FILE_FAIL" in l]
        assert fail_lines, "expected a READER_FILE_FAIL event"
        assert "detail=" in fail_lines[0], (
            f"detail field missing from log: {fail_lines[0]!r}"
        )
        assert "unparseable" in fail_lines[0], (
            f"detail should mention 'unparseable': {fail_lines[0]!r}"
        )


# ---------------------------------------------------------------------------
# verbatim_head: deterministic Python slice of extracted text
# ---------------------------------------------------------------------------


class TestVerbatimHead:
    def test_populated_with_lstripped_first_chars(
        self, workspace, fake_backend, logger
    ):
        """Reader stores extracted_text.lstrip()[:VERBATIM_HEAD_CHARS] on
        the catalog entry — leading whitespace removed, newlines preserved
        inside the slice."""
        from fda.organize import reader

        body = (
            "\n\n  \n"
            "Order ID: 10488\n"
            "\n"
            "Shipping Details:\n"
            "Ship Name: Frankenversand\n"
        )
        (workspace / "a.txt").write_text(body)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.verbatim_head.startswith("Order ID: 10488")
        assert "\n" in e.verbatim_head  # internal newlines preserved
        assert not e.verbatim_head.startswith("\n")
        assert not e.verbatim_head.startswith(" ")

    def test_capped_at_constant(self, workspace, fake_backend, logger):
        """The slice never exceeds VERBATIM_HEAD_CHARS characters."""
        from fda.organize import reader

        big = "x" * (reader.VERBATIM_HEAD_CHARS * 4)
        (workspace / "big.txt").write_text(big)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert len(e.verbatim_head) == reader.VERBATIM_HEAD_CHARS

    def test_cap_is_characters_not_bytes(self, workspace, fake_backend, logger):
        """Multibyte text must be capped at VERBATIM_HEAD_CHARS *characters*,
        not bytes. A regression to byte-based slicing (e.g. ``encode()[:N]``)
        would slice mid-codepoint and produce a shorter string on multibyte
        input — this test pins the character semantics."""
        from fda.organize import reader

        body = "界" * (reader.VERBATIM_HEAD_CHARS * 2)
        (workspace / "cjk.txt").write_text(body)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.verbatim_head == "界" * reader.VERBATIM_HEAD_CHARS

    def test_preserved_when_summary_call_fails(self, workspace, logger):
        """The slice is INDEPENDENT of the summarization call. When
        extraction succeeded but the backend errored out, verbatim_head is
        still the head of the extracted text — that grounding signal is
        the whole point of the field."""
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.side_effect = RuntimeError("boom")
        (workspace / "a.txt").write_text("Order ID: 10488\nShipping Details:\n")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.verbatim_head.startswith("Order ID: 10488")

    def test_preserved_when_summary_call_times_out(self, workspace, logger):
        """The TimeoutError branch in _summarize_one is distinct from the
        generic Exception branch. Pin that the slice survives it too —
        the docstring promises 'preserved across summarization timeout'."""
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.side_effect = TimeoutError("backend HTTP timeout")
        (workspace / "a.txt").write_text("Order ID: 10488\nShipping Details:\n")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.verbatim_head.startswith("Order ID: 10488")

    def test_preserved_when_summary_response_is_unparseable(
        self, workspace, logger
    ):
        """JSON parse failure path also preserves the slice."""
        from fda.organize import reader

        backend = MagicMock()
        backend.complete.return_value = "not-json"
        (workspace / "a.txt").write_text("Invoice\nOrder ID: 627\n")
        catalog = reader.read(workspace, backend=backend, logger=logger)
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.verbatim_head.startswith("Invoice")

    def test_empty_on_extractor_failure(self, workspace, fake_backend, logger):
        """When the EXTRACTOR returns a non-ok status (e.g., status='failed'
        because the file was unreadable or the extractor raised),
        verbatim_head is empty even though the helper sees an extraction
        result."""
        from fda.organize import reader, _extractors
        from fda.organize.models import ExtractionResult

        (workspace / "a.foo").write_bytes(b"\x00\x01\x02")

        def failing_extract(_path):
            return ExtractionResult(text=None, status="failed", note="boom")

        with patch.dict(_extractors.EXTRACTORS, {".foo": failing_extract}, clear=False):
            catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.extract_status == "failed"
        assert e.verbatim_head == ""

    def test_empty_for_unsupported_extension(self, workspace, fake_backend, logger):
        """No extractor → extract_status='no_extractor' → verbatim_head is ''."""
        from fda.organize import reader

        (workspace / "a.zzz").write_bytes(b"\x00\x01\x02\x03")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.extract_status == "no_extractor"
        assert e.verbatim_head == ""


# ---------------------------------------------------------------------------
# sections: copied from ExtractionResult through Reader to CatalogEntry,
# preserved across summarizer failure
# ---------------------------------------------------------------------------


class TestSectionsPropagation:
    def test_populated_from_extracted_text(
        self, workspace, fake_backend, logger
    ):
        """Reader copies extraction.sections into CatalogEntry.sections."""
        from fda.organize import reader

        body = (
            "Order ID: 10488\n"
            "\n"
            "Shipping Details:\n"
            "ACME Corp\n"
            "\n"
            "Customer Details:\n"
            "Hanna Moos\n"
        )
        (workspace / "a.txt").write_text(body)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.sections == ("Shipping Details", "Customer Details")

    def test_empty_when_extraction_returns_no_sections(
        self, workspace, fake_backend, logger
    ):
        """A JSON file (or any plaintext without colon-headers /
        ALL-CAPS dividers) yields sections=()."""
        from fda.organize import reader

        (workspace / "data.json").write_text(
            '{"customer_id": 1, "order_date": "2024-01-01"}\n'
        )
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = catalog.entries[0]
        assert e.sections == ()

    def test_preserved_through_summarizer_timeout(
        self, workspace, logger, monkeypatch
    ):
        """Mirrors the existing verbatim_head preservation contract:
        when the Haiku summarizer call times out, the deterministic
        sections list is still attached to the failed CatalogEntry."""
        from fda.organize import reader

        body = (
            "Shipping Details:\n"
            "ACME Corp\n"
            "\n"
            "Customer Details:\n"
            "Hanna Moos\n"
        )
        (workspace / "a.txt").write_text(body)

        class TimeoutBackend:
            def complete(self, **kwargs):
                raise TimeoutError("simulated backend timeout")

        catalog = reader.read(
            workspace, backend=TimeoutBackend(), logger=logger,
        )
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.sections == ("Shipping Details", "Customer Details")

    def test_preserved_through_summarizer_exception(
        self, workspace, logger
    ):
        from fda.organize import reader

        body = "Shipping Details:\nProduct XYZ\n"
        (workspace / "a.txt").write_text(body)

        class BoomBackend:
            def complete(self, **kwargs):
                raise RuntimeError("boom")

        catalog = reader.read(
            workspace, backend=BoomBackend(), logger=logger,
        )
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.sections == ("Shipping Details",)

    def test_preserved_through_unparseable_summary(
        self, workspace, logger
    ):
        from fda.organize import reader

        body = "Shipping Details:\nProduct XYZ\n"
        (workspace / "a.txt").write_text(body)

        class GarbageBackend:
            def complete(self, **kwargs):
                return "this is not JSON"

        catalog = reader.read(
            workspace, backend=GarbageBackend(), logger=logger,
        )
        e = catalog.entries[0]
        assert e.summary_failed is True
        assert e.sections == ("Shipping Details",)


class TestSectionsPropagationDocxXlsx:
    def test_docx_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader
        from docx import Document

        f = workspace / "doc.docx"
        d = Document()
        p = d.add_paragraph("My Title")
        p.style = d.styles["Title"]
        p2 = d.add_paragraph("Findings")
        p2.style = d.styles["Heading 1"]
        d.add_paragraph("body content")
        d.save(str(f))

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("doc.docx"))
        assert e.extract_status == "ok"
        assert e.sections == ("My Title", "Findings")
        assert e.verbatim_head.startswith("My Title")

    def test_xlsx_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader
        import openpyxl

        f = workspace / "wb.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders"
        ws.append(["Order ID", "Customer"])
        ws.append([1, "ACME"])
        wb.save(str(f))
        wb.close()

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("wb.xlsx"))
        assert e.extract_status == "ok"
        assert "Sheet:Orders" in e.sections
        assert "Order ID" in e.sections
        assert "Customer" in e.sections

    def test_docx_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Parallel to existing PDF/text coverage: when an extractor fails,
        the catalog entry's sections is preserved as ()."""
        from fda.organize import reader

        f = workspace / "broken.docx"
        f.write_bytes(b"not a real docx")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.docx"))
        assert e.extract_status == "failed"
        assert e.sections == ()

    def test_xlsx_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Spec line 166: reader failure preservation applies to both formats.
        Pin xlsx alongside docx."""
        from fda.organize import reader

        f = workspace / "broken.xlsx"
        f.write_bytes(b"not a real xlsx")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.xlsx"))
        assert e.extract_status == "failed"
        assert e.sections == ()

    def test_pptx_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader
        from pptx import Presentation

        f = workspace / "deck.pptx"
        prs = Presentation()
        for title in ("Quarter Plan", "Risks"):
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
        prs.save(str(f))

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("deck.pptx"))
        assert e.extract_status == "ok"
        assert e.sections == ("Quarter Plan", "Risks")
        assert e.verbatim_head.startswith("Slide 1: Quarter Plan")

    def test_pptx_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Parallel to docx/xlsx: when the extractor fails, the catalog
        entry's sections is preserved as ()."""
        from fda.organize import reader

        f = workspace / "broken.pptx"
        f.write_bytes(b"not a real pptx")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.pptx"))
        assert e.extract_status == "failed"
        assert e.sections == ()

    def test_csv_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader

        f = workspace / "data.csv"
        f.write_text(
            "customer_id,order_date,amount\n"
            "1,2024-01-01,100.00\n"
            "2,2024-01-02,200.00\n"
        )

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("data.csv"))
        assert e.extract_status == "ok"
        assert e.sections == ("customer_id", "order_date", "amount")
        # Verbatim head reflects the tab-normalized grid (delimiter normalized).
        assert e.verbatim_head.startswith("customer_id\torder_date\tamount")

    def test_csv_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Parallel to docx/xlsx/pptx: when the csv extractor fails (e.g. NUL
        byte triggers csv.Error mid-iteration), the catalog entry's sections
        is preserved as ()."""
        from fda.organize import reader

        f = workspace / "broken.csv"
        # NUL byte forces csv.reader to raise csv.Error('line contains NUL').
        f.write_bytes(b"name,age\nalice,30\n\x00bob,25\n")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.csv"))
        assert e.extract_status == "failed"
        assert e.sections == ()

    def test_hwpx_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader
        import zipfile

        ns = "http://www.hancom.co.kr/hwpml/2011/paragraph"
        f = workspace / "doc.hwpx"
        with zipfile.ZipFile(f, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            info = zipfile.ZipInfo("mimetype")
            info.compress_type = zipfile.ZIP_STORED
            zf.writestr(info, b"application/hwp+zip")
            xml = (
                f'<?xml version="1.0"?>'
                f'<hp:sec xmlns:hp="{ns}">'
                f'<hp:p><hp:run><hp:t>[발주서]</hp:t></hp:run></hp:p>'
                f'<hp:p><hp:run><hp:t>회사 정보:</hp:t></hp:run></hp:p>'
                f'</hp:sec>'
            ).encode("utf-8")
            zf.writestr("Contents/section0.xml", xml)

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("doc.hwpx"))
        assert e.extract_status == "ok"
        assert e.sections == ("발주서", "회사 정보")
        # Parallel to docx/pptx/csv reader tests — verbatim_head should hold.
        assert e.verbatim_head

    def test_hwpx_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Parallel to docx/xlsx/pptx/csv: when the hwpx extractor fails (not
        a zip), the catalog entry's sections is preserved as ()."""
        from fda.organize import reader

        f = workspace / "broken.hwpx"
        f.write_bytes(b"not a zip file")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.hwpx"))
        assert e.extract_status == "failed"
        assert e.sections == ()

    def test_hwpx_harvested_public_fixture_spot_check(
        self, workspace, fake_backend, logger, tmp_path
    ):
        """Spot-check against a real-world public-domain `.hwpx` (harvested in
        Task 0). Skips if the harvest gap was recorded — see spec spike findings."""
        from pathlib import Path
        import shutil
        from fda.organize import reader

        fixture_dir = Path(__file__).resolve().parent / "fixtures" / "hwp"
        harvested = sorted(fixture_dir.glob("*.hwpx"))
        if not harvested:
            pytest.skip("no harvested .hwpx fixtures (Task 0 gap)")

        # Copy the first harvested fixture into the workspace so reader sees it.
        sample = harvested[0]
        dest = workspace / sample.name
        shutil.copy2(sample, dest)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith(sample.name))
        # Real-world docs may or may not have headers _sections.py picks up;
        # the assertion is just that extraction succeeded with usable text.
        assert e.extract_status == "ok"
        # verbatim_head must be present so a fixture with non-standard
        # section paths doesn't silently pass with empty content.
        assert e.verbatim_head
        # Sections is a tuple — empty tuple is fine, but the type should hold.
        assert isinstance(e.sections, tuple)

    def test_hwp_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        """Real user sample round-trips through Reader; sections may be ()
        when the sample is PUA-heavy (documented v1 limitation)."""
        from pathlib import Path
        import shutil
        from fda.organize import reader

        samples_dir = Path(
            "/Users/hogyeongkim/Desktop/Projects/doc_agent_test_data/hwp_samples"
        )
        if not samples_dir.is_dir():
            pytest.skip(f"no .hwp samples at {samples_dir}")
        samples = sorted(samples_dir.glob("*.hwp"))
        if not samples:
            pytest.skip(f"no .hwp samples at {samples_dir}")

        sample = samples[0]
        dest = workspace / sample.name
        shutil.copy2(sample, dest)
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith(sample.name))
        assert e.extract_status == "ok"
        assert isinstance(e.sections, tuple)

    def test_hwp_failed_extraction_yields_empty_sections_in_catalog(
        self, workspace, fake_backend, logger
    ):
        """Non-OLE input → status="failed", sections=()."""
        from fda.organize import reader

        f = workspace / "broken.hwp"
        f.write_bytes(b"not an OLE compound document")
        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("broken.hwp"))
        assert e.extract_status == "failed"
        assert e.sections == ()


class TestKoreanSectionsThroughReader:
    """Korean structural-fingerprint coverage: a plaintext file containing
    Korean banners flows through Reader, and CatalogEntry.sections holds
    the expected Korean labels. Separate top-level class (not appended to
    TestSectionsPropagationDocxXlsx) so this PR doesn't collide with the
    in-flight CSV PR's edits to that class.
    """

    def test_korean_plaintext_sections_flow_through_reader(
        self, workspace, fake_backend, logger
    ):
        from fda.organize import reader

        f = workspace / "korean_doc.txt"
        f.write_text(
            "[발주서]\n"
            "회사 정보:\n"
            "■ 주의사항\n"
            "본문 내용...\n"
        )

        catalog = reader.read(workspace, backend=fake_backend, logger=logger)
        e = next(c for c in catalog.entries if c.path.endswith("korean_doc.txt"))
        assert e.extract_status == "ok"
        assert e.sections == ("발주서", "회사 정보", "주의사항")
