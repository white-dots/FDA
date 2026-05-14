# tests/test_organize_pipeline.py
"""End-to-end tests for fda.organize.organize() — Reader -> Classifier -> PlanBuilder -> Executor -> Verifier."""

from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from fda.organize import organize, apply_plan
from fda.organize.models import Plan, PlanResult


@pytest.fixture
def workspace(tmp_path):
    root = tmp_path / "ws"
    root.mkdir()
    (root / "a.txt").write_text("aaa")
    (root / "b.txt").write_text("bbb")
    (root / ".DS_Store").write_bytes(b"\x00")

    # Minimal .docx (one Title paragraph) so the integration corpus exercises
    # the new format end-to-end.
    from docx import Document
    docx_path = root / "report.docx"
    d = Document()
    p = d.add_paragraph("Report Title")
    p.style = d.styles["Title"]
    d.save(str(docx_path))

    # Minimal .xlsx (one sheet, one header row).
    import openpyxl
    xlsx_path = root / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Order ID", "Customer"])
    wb.save(str(xlsx_path))
    wb.close()

    # Minimal .pptx (one titled slide).
    from pptx import Presentation
    pptx_path = root / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Pipeline Demo"
    prs.save(str(pptx_path))

    # Minimal .csv (one header row, one data row).
    csv_path = root / "data.csv"
    csv_path.write_text("customer_id,order_date\n1,2024-01-01\n")

    # Minimal .hwpx (one section, one paragraph with a Korean bracket header).
    import zipfile as _zipfile
    hwpx_path = root / "korean.hwpx"
    ns = "http://www.hancom.co.kr/hwpml/2011/paragraph"
    with _zipfile.ZipFile(hwpx_path, "w", compression=_zipfile.ZIP_DEFLATED) as zf:
        info = _zipfile.ZipInfo("mimetype")
        info.compress_type = _zipfile.ZIP_STORED
        zf.writestr(info, b"application/hwp+zip")
        xml = (
            f'<?xml version="1.0"?>'
            f'<hp:sec xmlns:hp="{ns}">'
            f'<hp:p><hp:run><hp:t>[발주서]</hp:t></hp:run></hp:p>'
            f'</hp:sec>'
        ).encode("utf-8")
        zf.writestr("Contents/section0.xml", xml)

    return root


def _scripted_backend(workspace):
    """Parse the actual prompt body so the fake responds with real path_ids
    (path_ids are assigned in catalog order across ALL entries including
    junk, so '.DS_Store' shifts a.txt to f001 and b.txt to f002 — hardcoding
    f000/f001 would be wrong)."""
    backend = MagicMock()

    summary = json.dumps({"type_label": "text", "summary": "plain text"})
    taxonomy = json.dumps({
        "categories": [{
            "category_name": "Texts", "subpath": "Texts",
            "description": "plain text files", "criteria": "text-shaped",
        }],
        "fallback_category": {
            "category_name": "Misc", "subpath": "Misc",
            "description": "fallback", "criteria": "could not categorize",
        },
    })

    def fake_complete(*, messages, **_):
        body = messages[0]["content"]
        if "PATH:" in body:
            return summary
        # Stage A or Stage B is JSON
        try:
            parsed = json.loads(body)
        except json.JSONDecodeError:
            return ""
        if "BATCH" in parsed:
            return json.dumps({
                "assignments": [
                    {"path_id": e["path_id"], "category_name": "Texts"}
                    for e in parsed["BATCH"]
                ],
            })
        if "CATALOG" in parsed:
            return taxonomy
        return ""

    backend.complete.side_effect = fake_complete
    return backend


class TestOrganize:
    def test_full_pipeline(self, workspace):
        backend = _scripted_backend(workspace)
        result = organize(
            str(workspace),
            "sort",
            backend=backend,
            allowed_roots=[workspace.parent],
            metadata=False,
        )
        assert isinstance(result, PlanResult)
        # Both files moved into Texts/
        assert (workspace / "Texts" / "a.txt").exists()
        assert (workspace / "Texts" / "b.txt").exists()
        # Junk deleted
        assert not (workspace / ".DS_Store").exists()
        # New format files were also categorized (the scripted backend
        # assigns every entry to "Texts", so they all land in the same dir).
        assert (workspace / "Texts" / "report.docx").exists()
        assert (workspace / "Texts" / "data.xlsx").exists()
        assert (workspace / "Texts" / "deck.pptx").exists()
        assert (workspace / "Texts" / "data.csv").exists()
        assert (workspace / "Texts" / "korean.hwpx").exists()
        assert result.log_path is not None
        assert "Texts" in result.summary or "text-shaped" in result.summary

    def test_preview_returns_plan_without_executing(self, workspace):
        backend = _scripted_backend(workspace)
        plan = organize(
            str(workspace),
            "",
            preview=True,
            backend=backend,
            allowed_roots=[workspace.parent],
        )
        assert isinstance(plan, Plan)
        assert (workspace / "a.txt").exists()
        assert (workspace / ".DS_Store").exists()
        assert (workspace / "report.docx").exists()
        assert (workspace / "data.xlsx").exists()
        assert (workspace / "deck.pptx").exists()
        assert (workspace / "data.csv").exists()
        assert (workspace / "korean.hwpx").exists()
        assert plan.log_path is not None

    def test_invalid_target_raises(self, workspace):
        backend = MagicMock()
        with pytest.raises(ValueError):
            organize(
                str(workspace.parent),  # outside allowed_roots
                "",
                backend=backend,
                allowed_roots=[workspace],
            )


class TestOrganizeRouting:
    def test_routing_report_written_on_default_run(self, workspace):
        backend = _scripted_backend(workspace)
        # _scripted_backend handles summarizer / Stage A / Stage B; add a
        # router response by chaining the side_effect callable: any call
        # whose payload starts with '{"category"' is a router call.
        original = backend.complete.side_effect
        def _route_or_default(*args, **kwargs):
            body = kwargs["messages"][0]["content"]
            try:
                parsed = json.loads(body)
            except Exception:
                return original(*args, **kwargs)
            if "category" in parsed and "signals" in parsed:
                return json.dumps({
                    "destination": "sharepoint",
                    "reason": "test routing reason",
                    "misfits": [],
                })
            return original(*args, **kwargs)
        backend.complete.side_effect = _route_or_default

        # Routing-detection check is disjoint from Stage A/B: the
        # classifier's taxonomy payload contains the key `categories`
        # (plural), the assigner's payload contains `BATCH`, and the
        # router's payload contains both `category` (singular) AND
        # `signals` together. No Stage A/B call can satisfy that pair.
        organize(
            str(workspace), instructions="",
            backend=backend,
            allowed_roots=[workspace.parent],
            metadata=False,
        )
        assert (workspace / "routing-report.json").exists()
        assert (workspace / "routing-report.md").exists()
        parsed = json.loads((workspace / "routing-report.json").read_text())
        assert parsed["version"] == "1.0"

    def test_route_false_skips_router(self, workspace):
        backend = _scripted_backend(workspace)
        organize(
            str(workspace), instructions="",
            backend=backend,
            allowed_roots=[workspace.parent],
            route=False,
            metadata=False,
        )
        assert not (workspace / "routing-report.json").exists()
        assert not (workspace / "routing-report.md").exists()

    def test_preview_mode_skips_router(self, workspace):
        backend = _scripted_backend(workspace)
        result = organize(
            str(workspace), instructions="",
            backend=backend,
            allowed_roots=[workspace.parent],
            preview=True,
        )
        assert isinstance(result, Plan)
        assert not (workspace / "routing-report.json").exists()


class TestQuarantineEndToEnd:
    """End-to-end: mixed extractable + quarantine corpus lands the
    quarantine files under _NoExtractor/<ext>/ and the report shows
    them as Skipped sections."""

    @pytest.fixture
    def quarantine_only_workspace(self, tmp_path):
        """Workspace whose every file lacks a registered extractor.

        Deliberately does NOT use the shared `workspace` fixture (which
        creates extractable .txt/.docx/.xlsx/.pptx/.csv/.hwpx) — this
        test needs the classifier's empty-extractable early-return path.
        """
        root = tmp_path / "ws-q"
        root.mkdir()
        (root / "a.doc").write_bytes(b"a")
        (root / "b.xls").write_bytes(b"b")
        return root

    def test_unsupported_extension_moved_to_NoExtractor_bucket(self, workspace):
        from fda.organize import organize

        # `workspace` fixture already created extractable files (.txt, .docx,
        # .xlsx, .pptx, .csv, .hwpx). Add one unsupported-extension file so
        # the test exercises both paths in the same run.
        (workspace / "old-quote.doc").write_bytes(b"fake doc bytes")
        backend = _scripted_backend(workspace)
        # Router is invoked because route=True (default) — chain a router
        # response onto the scripted backend, mirroring how TestOrganizeRouting
        # at test_organize_pipeline.py:178-192 does it.
        original = backend.complete.side_effect
        def _route_or_default(*args, **kwargs):
            body = kwargs["messages"][0]["content"]
            try:
                parsed = json.loads(body)
            except Exception:
                return original(*args, **kwargs)
            if "category" in parsed and "signals" in parsed:
                return json.dumps({
                    "destination": "sharepoint",
                    "reason": "test routing reason",
                    "misfits": [],
                })
            return original(*args, **kwargs)
        backend.complete.side_effect = _route_or_default

        organize(
            str(workspace),
            instructions="",
            backend=backend,
            allowed_roots=[workspace.parent],
        )
        # The .doc goes to quarantine; extractable files keep flowing through
        # their normal categories (assertions for those already live in
        # TestOrganize.test_full_pipeline above).
        assert (workspace / "_NoExtractor" / "doc" / "old-quote.doc").exists()
        md = (workspace / "routing-report.md").read_text(encoding="utf-8")
        assert "## 건너뜀 — 추출기 없음 (No Extractor)" in md
        assert "`_NoExtractor/doc/old-quote.doc`" in md
        data = json.loads(
            (workspace / "routing-report.json").read_text(encoding="utf-8"),
        )
        assert any(g["bucket"] == "_NoExtractor" and g["ext"] == "doc"
                   for g in data["quarantine"])

    def test_all_quarantine_corpus_does_not_crash(self, quarantine_only_workspace):
        """No extractable files at all → classifier short-circuits, plan_builder
        emits only quarantine MOVEs, router emits only Skipped sections."""
        from fda.organize import organize

        workspace = quarantine_only_workspace
        # The classifier short-circuits before any LLM call, so the scripted
        # backend never sees a summarizer/Stage A/Stage B payload. A bare
        # MagicMock with a never-called .complete is enough here.
        backend = MagicMock()
        organize(
            str(workspace),
            instructions="",
            backend=backend,
            allowed_roots=[workspace.parent],
        )
        assert backend.complete.call_count == 0
        assert (workspace / "_NoExtractor" / "doc" / "a.doc").exists()
        assert (workspace / "_NoExtractor" / "xls" / "b.xls").exists()
        md = (workspace / "routing-report.md").read_text(encoding="utf-8")
        assert "건너뜀: 2 (추출기 없음 2, 추출 실패 0)" in md


def test_translate_catalog_for_stage5_forwards_files_pinned():
    from fda.organize import _translate_catalog_for_stage5
    from fda.organize.models import Catalog
    catalog = Catalog(
        target="/tmp/x",
        entries=(),
        git_repos_skipped=("/tmp/x/.git",),
        files_pinned=("/tmp/x/manifest.csv", "/tmp/x/README.md"),
    )
    translated = _translate_catalog_for_stage5(catalog, outcomes=())
    assert translated.files_pinned == (
        "/tmp/x/manifest.csv", "/tmp/x/README.md",
    )
    assert translated.git_repos_skipped == ("/tmp/x/.git",)


def test_organize_leaves_manifest_at_root(workspace, tmp_path):
    """manifest.csv at the root survives organize() untouched."""
    (workspace / "manifest.csv").write_text("id,name\n1,a\n")
    backend = _scripted_backend(workspace)
    result = organize(
        str(workspace), instructions="", backend=backend,
        allowed_roots=[tmp_path],
        route=False, metadata=False,
    )
    assert (workspace / "manifest.csv").is_file()
    assert (workspace / "manifest.csv").read_text() == "id,name\n1,a\n"


def test_organize_no_ghost_folders_when_pinned(workspace, tmp_path):
    """A pinned root file does not anchor a ghost subfolder.

    Setup: place an organizable file in a nested folder (`inbox/`) that the
    organize pipeline should drain. After the run, the source folder must be
    gone (verifier rmdir'd it) AND root pinned files must still be there.
    """
    (workspace / "manifest.csv").write_text("id\n1\n")
    (workspace / "README.md").write_text("# x\n")
    inbox = workspace / "inbox"
    inbox.mkdir()
    (inbox / "inner.txt").write_text("inner")  # this file should be organized away
    backend = _scripted_backend(workspace)
    result = organize(
        str(workspace), instructions="", backend=backend,
        allowed_roots=[tmp_path],
        route=False, metadata=False,
    )
    # Verifier cleaned the emptied inbox folder (it appears in leftover_empty_dirs
    # because that field tracks dirs the verifier successfully rmdir'd).
    assert str(inbox) in result.leftover_empty_dirs
    # Emptied source folder is gone.
    assert not inbox.exists()
    # Pinned files still at root.
    assert (workspace / "manifest.csv").is_file()
    assert (workspace / "README.md").is_file()


def test_pinned_takes_precedence_over_quarantine(workspace, tmp_path):
    """A pinned filename with an unextractable extension is pinned, not quarantined."""
    (workspace / ".fda-ignore").write_text("inventory.xyz\n")
    (workspace / "inventory.xyz").write_bytes(b"\x00\x01\x02")  # unextractable ext
    backend = _scripted_backend(workspace)
    result = organize(
        str(workspace), instructions="", backend=backend,
        allowed_roots=[tmp_path],
        route=False, metadata=False,
    )
    # File is still at root, was never moved into a quarantine bucket.
    assert (workspace / "inventory.xyz").is_file()
    assert not (workspace / "_NoExtractor").exists()
    assert not (workspace / "_ExtractionFailed").exists()
